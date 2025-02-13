from abc import ABC, abstractmethod
from docx import Document
import openpyxl


from kivy.app import App
from kivy.lang import Builder
from kivy.uix.boxlayout import BoxLayout
from kivy.uix.scrollview import ScrollView
from kivy.uix.label import Label
from kivy.uix.button import Button
from kivy.uix.textinput import TextInput
from kivy.uix.spinner import Spinner


class BankService(ABC):
    def __init__(self, principal=0.0, rate=0.0, term=0):
        self._principal = principal
        self._rate = rate
        self._term = term


    @property
    def principal(self):
        return self._principal


    @principal.setter
    def principal(self, value):
        self._principal = value


    @property
    def rate(self):
        return self._rate


    @rate.setter
    def rate(self, value):
        self._rate = value


    @property
    def term(self):
        return self._term


    @term.setter
    def term(self, value):
        self._term = value


    @abstractmethod
    def calculate(self):
        """Вычисление платежного графика или накоплений"""
        pass


class CreditService(BankService):
    def calculate(self):
        """Расчёт аннуитетного кредита"""
        schedule = []
        monthly_rate = self.rate / 12 / 100
        annuity = (self.principal * monthly_rate) / (1 - (1 + monthly_rate) ** (-self.term))
        balance = self.principal
        for month in range(1, self.term + 1):
            interest = balance * monthly_rate
            principal_payment = annuity - interest
            balance -= principal_payment
            schedule.append((month, round(annuity, 2), round(balance, 2)))
        return schedule


    def __str__(self):
        return f"CreditService(principal={self.principal}, rate={self.rate}, term={self.term})"


    def __repr__(self):
        return self.__str__()


class InstallmentService(BankService):
    def calculate(self):
        """Расчёт рассрочки без процентов"""
        schedule = []
        monthly_payment = self.principal / self.term
        balance = self.principal
        for month in range(1, self.term + 1):
            balance -= monthly_payment
            schedule.append((month, round(monthly_payment, 2), round(balance, 2)))
        return schedule


    def __str__(self):
        return f"InstallmentService(principal={self.principal}, term={self.term})"


    def __repr__(self):
        return self.__str__()


class DepositService(BankService):
    def calculate(self):
        """Расчёт вклада с ежемесячной капитализацией процентов"""
        schedule = []
        balance = self.principal
        monthly_rate = self.rate / 12 / 100
        for month in range(1, self.term + 1):
            balance += balance * monthly_rate
            schedule.append((month, round(self.rate, 2), round(balance, 2)))
        return schedule


    def __str__(self):
        return f"DepositService(principal={self.principal}, rate={self.rate}, term={self.term})"


    def __repr__(self):
        return self.__str__()


class BankAppGUI(BoxLayout):
    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.orientation = 'vertical'
        self.padding = 10
        self.spacing = 10


        self.add_widget(Label(text='Выберите услугу:', size_hint=(1, None), height=30))
        self.service_spinner = Spinner(
            text='Кредит',
            values=('Кредит', 'Рассрочка', 'Вклад'),
            size_hint=(1, None),
            height=44
        )
        self.add_widget(self.service_spinner)


        self.principal_input = TextInput(hint_text='Сумма', multiline=False, input_filter='float', size_hint=(1, None), height=44)
        self.rate_input = TextInput(hint_text='Ставка (%)', multiline=False, input_filter='float', size_hint=(1, None), height=44)
        self.term_input = TextInput(hint_text='Срок (мес.)', multiline=False, input_filter='int', size_hint=(1, None), height=44)
        self.add_widget(self.principal_input)
        self.add_widget(self.rate_input)
        self.add_widget(self.term_input)


        self.calc_button = Button(text='Рассчитать', size_hint=(1, None), height=44)
        self.calc_button.bind(on_press=self.perform_calculation)
        self.add_widget(self.calc_button)


        self.result_label = Label(text='', size_hint_y=None, markup=True)
        self.result_label.bind(texture_size=self.result_label.setter('size'))
        result_scroll = ScrollView(size_hint=(1, 1))
        result_scroll.add_widget(self.result_label)
        self.add_widget(result_scroll)


        self.save_docx_button = Button(text='Сохранить как DOCX', size_hint=(1, None), height=44)
        self.save_docx_button.bind(on_press=self.save_docx)
        self.add_widget(self.save_docx_button)


        self.save_xlsx_button = Button(text='Сохранить как XLSX', size_hint=(1, None), height=44)
        self.save_xlsx_button.bind(on_press=self.save_xlsx)
        self.add_widget(self.save_xlsx_button)


        self.current_result = None


    def perform_calculation(self, instance):
        try:
            principal = float(self.principal_input.text)
            rate = float(self.rate_input.text)
            term = int(self.term_input.text)
            service_type = self.service_spinner.text


            if service_type == 'Кредит':
                service = CreditService(principal, rate, term)
            elif service_type == 'Рассрочка':
                service = InstallmentService(principal, 0, term)
            elif service_type == 'Вклад':
                service = DepositService(principal, rate, term)
            else:
                self.result_label.text = "Неизвестный тип услуги."
                return


            result = service.calculate()
            self.current_result = result
            result_str = f"[b]{service}[/b]\n"
            for row in result:
                result_str += str(row) + "\n"
            self.result_label.text = result_str


        except Exception as e:
            self.result_label.text = f"Ошибка: {str(e)}"


    def save_docx(self, instance):
        if not self.current_result:
            self.result_label.text = "Сначала выполните расчёт."
            return
        try:
            doc = Document()
            doc.add_paragraph(self.result_label.text)
            doc.save("report.docx")
            self.result_label.text += "\nРезультаты сохранены в report.docx"
        except Exception as e:
            self.result_label.text += f"\nОшибка при сохранении DOCX: {str(e)}"


    def save_xlsx(self, instance):
        if not self.current_result:
            self.result_label.text = "Сначала выполните расчёт."
            return
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["Месяц", "Платёж/Ставка", "Остаток/Сумма"])
            for row in self.current_result:
                ws.append(row)
            wb.save("report.xlsx")
            self.result_label.text += "\nРезультаты сохранены в report.xlsx"
        except Exception as e:
            self.result_label.text += f"\nОшибка при сохранении XLSX: {str(e)}"


class BankApp(App):
    def build(self):
        return BankAppGUI()


if __name__ == '__main__':
    BankApp().run()
