import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from docx import Document
import openpyxl

import bank_services.credit as credit
import bank_services.installment as installment
import bank_services.deposit as deposit

class BankApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Банковские услуги")
        self.root.geometry("400x600")

        self.service_var = tk.StringVar(value="Кредит")
        ttk.Label(root, text="Выберите услугу:").pack(pady=5)
        self.service_menu = ttk.Combobox(root, textvariable=self.service_var, 
                                         values=["Кредит", "Рассрочка", "Вклад"])
        self.service_menu.pack(pady=5)
        self.service_menu.bind("<<ComboboxSelected>>", self.update_fields)

        self.param1_label = ttk.Label(root, text="Сумма:")
        self.param1_label.pack(pady=5)
        self.param1_entry = ttk.Entry(root)
        self.param1_entry.pack(pady=5)

        self.param2_label = ttk.Label(root, text="Ставка (%):")
        self.param2_label.pack(pady=5)
        self.param2_entry = ttk.Entry(root)
        self.param2_entry.pack(pady=5)

        self.param3_label = ttk.Label(root, text="Срок (мес.):")
        self.param3_label.pack(pady=5)
        self.param3_entry = ttk.Entry(root)
        self.param3_entry.pack(pady=5)

        self.rate_type_var = tk.StringVar(value="Фиксированная")
        ttk.Label(root, text="Тип ставки:").pack(pady=5)
        self.rate_fixed = ttk.Radiobutton(root, text="Фиксированная",
                                          variable=self.rate_type_var,
                                          value="Фиксированная",
                                          command=self.toggle_rate_input)
        self.rate_fixed.pack()
        self.rate_variable = ttk.Radiobutton(root, text="Изменяемая",
                                             variable=self.rate_type_var,
                                             value="Изменяемая",
                                             command=self.toggle_rate_input)
        self.rate_variable.pack()

        self.custom_rates_entry = tk.Text(root, height=3, width=40)
        self.custom_rates_entry.pack_forget()

        self.calc_button = ttk.Button(root, text="Рассчитать", command=self.calculate)
        self.calc_button.pack(pady=10)

        self.result_text = tk.Text(root, height=10, width=50)
        self.result_text.pack(pady=5)

        self.save_button = ttk.Button(root, text="Сохранить", command=self.save_results)
        self.save_button.pack(pady=10)

    def update_fields(self, event):
        service = self.service_var.get()
        if service == "Рассрочка":
            self.param2_label.pack_forget()
            self.param2_entry.pack_forget()
        else:
            self.param2_label.pack(pady=5)
            self.param2_entry.pack(pady=5)

    def toggle_rate_input(self):
        """
        Показывает или скрывает поле для ввода изменяемых ставок в зависимости от выбранного типа ставки.
        """
        if self.rate_type_var.get() == "Изменяемая":
            self.custom_rates_entry.pack(pady=5)
        else:
            self.custom_rates_entry.pack_forget()

    def calculate(self):
        try:
            principal = float(self.param1_entry.get())
            term = int(self.param3_entry.get())
            service = self.service_var.get()
            result = []

            if service == "Кредит":
                rate = float(self.param2_entry.get())
                result = credit.calculate_credit(principal, rate, term)
            elif service == "Рассрочка":
                result = installment.calculate_installment(principal, term)
            elif service == "Вклад":
                if self.rate_type_var.get() == "Фиксированная":
                    rate = float(self.param2_entry.get())
                    result = deposit.calculate_deposit(principal, rate, term)
                else:
                    rates_text = self.custom_rates_entry.get("1.0", tk.END).strip()
                    if not rates_text:
                        messagebox.showerror("Ошибка", "Введите ставки для каждого месяца через запятую!")
                        return
                    try:
                        custom_rates = list(map(float, rates_text.split(",")))
                        if len(custom_rates) != term:
                            messagebox.showerror("Ошибка", f"Введите {term} значений ставок!")
                            return
                        result = deposit.calculate_deposit(principal, 0, term, custom_rates)
                    except ValueError:
                        messagebox.showerror("Ошибка", "Некорректный формат ставок!")
                        return

            self.result_text.delete("1.0", tk.END)
            for row in result:
                self.result_text.insert(tk.END, f"{row}\n")

            self.calculation_result = result

        except ValueError:
            messagebox.showerror("Ошибка", "Введите корректные данные!")

    def save_results(self):
        filetypes = [("Документ Word", "*.docx"), ("Таблица Excel", "*.xlsx")]
        filepath = filedialog.asksaveasfilename(filetypes=filetypes, defaultextension=filetypes)

        if not filepath:
            return

        if filepath.endswith(".docx"):
            doc = Document()
            doc.add_paragraph(str(self.calculation_result))
            doc.save(filepath)
        elif filepath.endswith(".xlsx"):
            wb = openpyxl.Workbook()
            ws = wb.active
            for row in self.calculation_result:
                ws.append(row)
            wb.save(filepath)

if __name__ == "__main__":
    root = tk.Tk()
    app = BankApp(root)
    root.mainloop()